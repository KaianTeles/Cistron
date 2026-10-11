#!/usr/bin/env python3
"""Adversarial test suite for cistron."""
import os, sys, io, itertools, traceback, hashlib
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import xml.etree.ElementTree as ET
from openpyxl import Workbook
import cistron as M

PASS, FAIL, WARN = [], [], []
def ok(n, d=""):   PASS.append((n, d))
def bad(n, d=""):  FAIL.append((n, d))
def soft(n, d=""): WARN.append((n, d))

# scratch goes to the system temp dir, never into the working tree
import atexit, shutil, tempfile
TMP = tempfile.mkdtemp(prefix="cistron_tests_")
atexit.register(shutil.rmtree, TMP, True)

def make_wb(path, cassettes, settings=None, colors=None, sheets=("settings","colors","cassettes")):
    wb = Workbook(); wb.remove(wb.active)
    if "settings" in sheets:
        ws = wb.create_sheet("settings"); ws.append(["setting","value"])
        for k, v in (settings or {}).items(): ws.append([k, v])
    if "colors" in sheets:
        ws = wb.create_sheet("colors"); ws.append(["name","hex"])
        for k, v in (colors or {}).items(): ws.append([k, v])
    if "cassettes" in sheets:
        ws = wb.create_sheet("cassettes")
        ws.append(["cassette","order","type","label","size","dir","join","color"])
        for r in cassettes: ws.append(list(r) + [None]*(8-len(r)))
    wb.save(path); return path

SIMPLE = [("A",1,"HA","Left","M","",""),("A",2,"prom","P1","","",""),
          ("A",3,"cds","Gene","M","","yes"),("A",4,"term","","","",""),
          ("A",5,"HA","Right","M","","")]

def render(path, **over):
    s, c, cons = M.load_workbook_config(path)
    s.update(over)
    st, pal, _, _ = M.resolve_theme(s)
    res = M.assign_palette(cons, c, pal, str(s["palette_mode"]).lower())
    return M.build_figure(cons, s, res, st), s, cons, st

def svg_text(path):
    return open(path).read()

# ---------------------------------------------------------------- 1. errors
print("== malformed input ==")
cases = [
    ("no cassettes sheet", dict(cassettes=[], sheets=("settings","colors")), "no 'cassettes' sheet"),
    ("empty cassettes",    dict(cassettes=[]), "no usable rows"),
    ("unknown type",       dict(cassettes=[("A",1,"wormhole","X","M","","")]), "unknown type"),
    ("bad size",           dict(cassettes=[("A",1,"cds","X","XXL","","")]), "size must be"),
    ("bad order",          dict(cassettes=[("A","first","cds","X","M","","")]), "order must be"),
]
for name, kw, expect in cases:
    p = make_wb(f"{TMP}/e_{abs(hash(name))}.xlsx", **kw)
    try:
        render(p); bad(name, "no error raised")
    except SystemExit as e:
        (ok if expect in str(e) else bad)(name, str(e).split("\n")[0][:70])
    except Exception as e:
        bad(name, f"{type(e).__name__}: {e}")

for name, over, expect in [
    ("unknown theme",   dict(theme="bogus"),        "Unknown theme"),
    ("bad anchor at render", dict(anchor_side="middle"), "anchor_side must"),
    ("unknown style",   dict(style="bogus"),        "Unknown style"),
    ("unknown palette", dict(palette="bogus"),      "Unknown palette"),
    ("bad labels",      dict(labels="sideways"),    "labels must be"),
    ("bad repeat",      dict(label_repeat="twice"), "label_repeat must"),
]:
    p = make_wb(f"{TMP}/s_{abs(hash(name))}.xlsx", SIMPLE)
    try:
        s, c, cons = M.load_workbook_config(p); s.update(over)
        st, pal, _, _ = M.resolve_theme(s)
        res = M.assign_palette(cons, c, pal, str(s["palette_mode"]).lower())
        M.build_figure(cons, s, res, st)
        bad(name, "no error raised")
    except SystemExit as e:
        (ok if expect in str(e) else bad)(name, str(e)[:70])
    except Exception as e:
        bad(name, f"{type(e).__name__}: {e}")

# settings validated at load time, not render time
for name, setting, expect in [
    ("bad palette_mode at load", {"palette_mode":"bogus"}, "palette_mode must"),
    ("bad anchor at load",       {"anchor_side":"middle"}, "anchor_side must"),
]:
    p = make_wb(f"{TMP}/l_{abs(hash(name))}.xlsx", SIMPLE, settings=setting)
    try:
        M.load_workbook_config(p); bad(name, "no error raised")
    except SystemExit as e:
        (ok if expect in str(e) else bad)(name, str(e)[:60])

p = make_wb(f"{TMP}/missing.xlsx", SIMPLE)
try:
    M.load_workbook_config(f"{TMP}/does_not_exist.xlsx"); bad("missing file")
except Exception as e:
    ok("missing file", type(e).__name__)

# ---------------------------------------------------------- 2. odd but legal
print("== odd but legal ==")
odd = {
 "single element":    [("A",1,"cds","Solo","M","","")],
 "join on first":     [("A",1,"cds","First","M","","yes"),("A",2,"cds","Second","M","","")],
 "blank order":       [("A",None,"HA","L","M","",""),("A",None,"cds","G","M","","")],
 "duplicate order":   [("A",1,"HA","L","M","",""),("A",1,"cds","G","M","","")],
 "no labels at all":  [("A",1,"HA",None,"M","",""),("A",2,"cds",None,"M","","")],
 "xml-hostile label": [("A",1,"cds","p<35> & \"Q\" 'x'","M","",""),("A",2,"HA","R","M","","")],
 "unicode label":     [("A",1,"cds","αβγ ΔpsbA→ 中文","M","",""),("A",2,"HA","R","M","","")],
 "very long label":   [("A",1,"cds","A"*120,"M","",""),("A",2,"HA","R","M","","")],
 "all reversed":      [("A",i+1,t,f"X{i}","M","rev","") for i,t in
                       enumerate(["HA","prom","cds","ctp","tag","p2a","term","loxp"])],
 "every type":        [("A",i+1,t,t,"M","","") for i,t in enumerate(sorted(M.PART_TYPES))],
 "many cassettes":    [(f"C{j}",i+1,"cds",f"g{i}","S","","") for j in range(25) for i in range(3)],
 "lowercase types":   [("A",1,"ha","L","M","",""),("A",2,"GENE","G","M","","")],
 "spacey values":     [("  A  ",1,"  cds  ","  G  ","  m  ","  REV  ","  YES  ")],
 "hex without hash":  [("A",1,"cds","G","M","","","2C7FB8"),("A",2,"HA","R","M","","")],
}
for name, rows in odd.items():
    p = make_wb(f"{TMP}/o_{abs(hash(name))}.xlsx", rows)
    try:
        (out, w, h, fpx, lh, pr), s, cons, st = render(p, labels="outside", figure_width_cm=18)
        txt = svg_text(out)
        ET.fromstring(txt)                       # must be well-formed XML
        if "nan" in txt.lower() or "inf" in txt.lower():
            bad(name, "NaN/Inf in output")
        elif w <= 0 or h <= 0:
            bad(name, f"bad viewBox {w}x{h}")
        else:
            ok(name, f"{w:.0f}x{h:.0f}")
    except Exception as e:
        bad(name, f"{type(e).__name__}: {e}")

# ------------------------------------------------------- 3. extreme settings
print("== extreme settings ==")
p = make_wb(f"{TMP}/ext.xlsx", SIMPLE)
for name, over in {
 "width 0.5cm":       dict(figure_width_cm=0.5),
 "width 200cm":       dict(figure_width_cm=200),
 "height 0.1cm":      dict(figure_width_cm=10, figure_height_cm=0.1),
 "label_pt 1":        dict(labels="inside", label_pt=1),
 "label_pt 80":       dict(labels="inside", label_pt=80),
 "spacing 0":         dict(cassette_spacing=0, labels="outside"),
 "spacing 2000":      dict(cassette_spacing=2000),
 "line_thickness 0":  dict(line_thickness=0),
 "line_thickness 30": dict(line_thickness=30),
 "min_font_pt 60":    dict(min_font_pt=60, show_legend="TRUE"),
 "no width":          dict(figure_width_cm=None),
}.items():
    try:
        (out, w, h, fpx, lh, pr), s, cons, st = render(p, **over)
        ET.fromstring(svg_text(out))
        if w <= 0 or h <= 0: bad(name, f"viewBox {w}x{h}")
        else: ok(name, f"{w:.0f}x{h:.0f}")
    except Exception as e:
        bad(name, f"{type(e).__name__}: {e}")

# ------------------------------------------- 4. full theme x label matrix
print("== theme x labels x repeat matrix ==")
p = make_wb(f"{TMP}/mat.xlsx", [
    ("A",1,"HA","Left arm","M","",""),("A",2,"prom","Promoter","","",""),
    ("A",3,"ctp","CTP","XS","","yes"),("A",4,"cds","Reporter","L","","yes"),
    ("A",5,"tag","Tag","XS","","yes"),("A",6,"term","","","",""),
    ("A",7,"loxp","Site","XS","",""),("A",8,"cds","Marker","M","rev",""),
    ("A",9,"HA","Right arm","M","",""),
    ("B",1,"HA","Left arm","M","",""),("B",2,"prom","Promoter","","",""),
    ("B",3,"cds","Other","L","",""),("B",4,"HA","Right arm","M","",""),
])
fails = 0; n = 0
for theme, lab, rep in itertools.product(sorted(M.THEMES),
                                         ["none","inside","outside","below"],
                                         ["all","once","row"]):
    n += 1
    try:
        (out, w, h, fpx, lh, pr), s, cons, st = render(
            p, theme=theme, style="", palette="", palette_mode="theme",
            labels=lab, label_repeat=rep, figure_width_cm=18,
            outfile=f"{TMP}/m.svg")
        t = svg_text(out); ET.fromstring(t)
        if "nan" in t.lower(): raise ValueError("NaN")
    except Exception as e:
        fails += 1; bad(f"{theme}/{lab}/{rep}", f"{type(e).__name__}: {e}")
(ok if fails == 0 else bad)(f"matrix {n} combinations", f"{n-fails} rendered, {fails} failed")

# style x palette grid
fails = 0; n = 0
for sty, pal in itertools.product(sorted(M.STYLES), sorted(M.PALETTES)):
    n += 1
    try:
        render(p, theme="", style=sty, palette=pal, palette_mode="theme",
               labels="inside", figure_width_cm=18, outfile=f"{TMP}/g.svg")
    except Exception as e:
        fails += 1; bad(f"{sty}+{pal}", str(e)[:50])
(ok if fails == 0 else bad)(f"grid {n} style x palette", f"{n-fails} ok")

# ------------------------------------------------------------- 5. invariants
print("== invariants ==")
# determinism
h1 = h2 = None
for i in range(2):
    (out, *_), *_ = render(p, theme="journal", labels="outside",
                           figure_width_cm=16, outfile=f"{TMP}/det{i}.svg")
    d = hashlib.sha1(open(out,"rb").read()).hexdigest()
    h1, h2 = (d, h2) if i == 0 else (h1, d)
(ok if h1 == h2 else bad)("deterministic output", h1[:10])

# labels=none must not change geometry vs the pre-label behaviour
(a, wa, ha_, *_), *_ = render(p, labels="none", show_legend="FALSE",
                                figure_width_cm=16, outfile=f"{TMP}/n1.svg")
(b, wb_, hb, *_), *_ = render(p, labels="none", label_repeat="once",
                                show_legend="FALSE", figure_width_cm=16,
                                outfile=f"{TMP}/n2.svg")
(ok if (wa, ha_) == (wb_, hb) else bad)("repeat is a no-op when labels=none",
                                        f"{wa}x{ha_} vs {wb_}x{hb}")

# once <= all, in label count
import re
def ntext(path): return len(re.findall(r"<text", svg_text(path)))
(x, *_), *_ = render(p, labels="outside", label_repeat="all",
                     figure_width_cm=16, outfile=f"{TMP}/r1.svg")
(y, *_), *_ = render(p, labels="outside", label_repeat="once",
                     figure_width_cm=16, outfile=f"{TMP}/r2.svg")
na, no = ntext(x), ntext(y)
(ok if no < na else bad)("once prints fewer labels", f"{na} -> {no}")

# inside labels must not change width when suppressed
(x, w1, *_), *_ = render(p, labels="inside", label_repeat="all",
                         figure_width_cm=16, outfile=f"{TMP}/i1.svg")
(y, w2, *_), *_ = render(p, labels="inside", label_repeat="once",
                         figure_width_cm=16, outfile=f"{TMP}/i2.svg")
(ok if w1 == w2 else bad)("suppressed inside label keeps width", f"{w1} vs {w2}")

# anchor side actually differs
(x, *_), *_ = render(p, anchor_side="left",  figure_width_cm=16, outfile=f"{TMP}/al.svg")
(y, *_), *_ = render(p, anchor_side="right", figure_width_cm=16, outfile=f"{TMP}/ar.svg")
(ok if svg_text(x) != svg_text(y) else bad)("anchor_side changes layout")

# runaway guard
for name, over, limit in [("tiny width guarded", dict(figure_width_cm=0.5), 20000),
                          ("huge min_font guarded", dict(figure_width_cm=8, min_font_pt=60,
                                                         show_legend="TRUE"), 20000)]:
    (out, w, h, fpx, lh, pr), s_, cons_, st_ = render(p, outfile=f"{TMP}/rg.svg", **over)
    wn = M.fit_warnings(cons_, s_, st_, w, h, fpx, lh, pr)
    hit = any("unreachable" in x for x in wn)
    (ok if (w < limit and h < limit and hit) else bad)(
        name, f"{w:.0f}x{h:.0f}, warned={hit}")

# blank cassette cells continue the group above (and must not merge groups)
pff = make_wb(f"{TMP}/ff.xlsx", [
    ("Construct A",1,"HA","L","M","",""), (None,2,"prom","P1","","",""),
    (None,3,"cds","GeneA","M","",""),
    ("Construct B",1,"HA","L","M","",""), (None,2,"prom","P2","","",""),
    (None,3,"cds","GeneB","M","","")])
pex = make_wb(f"{TMP}/ff2.xlsx", [
    ("Construct A",1,"HA","L","M","",""), ("Construct A",2,"prom","P1","","",""),
    ("Construct A",3,"cds","GeneA","M","",""),
    ("Construct B",1,"HA","L","M","",""), ("Construct B",2,"prom","P2","","",""),
    ("Construct B",3,"cds","GeneB","M","","")])
def shape(path):
    _, _, cs = M.load_workbook_config(path)
    return [(c["name"], [p["label"] for p in c["parts"]]) for c in cs]
sf, se = shape(pff), shape(pex)
(ok if len(sf) == 2 else bad)("blank cassette cells do not merge groups", f"{len(sf)} cassettes")
(ok if sf == se else bad)("blank-fill matches a fully repeated column")

# ------------------------------------------------------------- 6. contrast
print("== contrast ==")
worst = []
for theme in sorted(M.THEMES):
    s, c, cons = M.load_workbook_config(p)
    s.update(theme=theme, style="", palette="", palette_mode="theme", labels="inside")
    st, pal, _, _ = M.resolve_theme(s)
    st = dict(st); st["_sw"] = s["line_thickness"] * st["stroke_w"]
    res = M.assign_palette(cons, c, pal, "theme")
    for key, col in res.items():
        f, stroke, sw = M.ink(col, st)
        # on hollow styles the fill IS the background on purpose; the
        # outline is what carries the colour, so measure that instead
        carrier = f if st["fill"] == "solid" else (stroke if stroke != "none" else f)
        d_bg = abs(M.luminance(carrier) - M.luminance(st["bg"]))
        txt = M.label_ink(col, st)
        d_tx = abs(M.luminance(f) - M.luminance(txt))
        worst.append((d_bg, d_tx, theme, key[1] or key[0]))
glyph_bad = [w for w in worst if w[0] < 0.045]
text_bad  = [w for w in worst if w[1] < 0.30]
(ok if not glyph_bad else soft)("glyph vs background",
    f"min delta-L {min(w[0] for w in worst):.3f}" +
    (f" — {len(glyph_bad)} under 0.045, e.g. {glyph_bad[0][2]}/{glyph_bad[0][3]}" if glyph_bad else ""))
(ok if not text_bad else soft)("inside text vs fill",
    f"min delta-L {min(w[1] for w in worst):.3f}" +
    (f" — {len(text_bad)} under 0.30, e.g. {text_bad[0][2]}/{text_bad[0][3]}" if text_bad else ""))

# ------------------------------------------------------------- 7. warnings
print("== sanity-check behaviour ==")
(out, w, h, fpx, lh, pr), s, cons, st = render(p, labels="outside", label_pt=11,
                                           figure_width_cm=16, outfile=f"{TMP}/w.svg")
wn = M.fit_warnings(cons, s, st, w, h, fpx, lh, pr)
(ok if not wn else bad)("no false positive on sane settings", f"{len(wn)} warnings")
(out, w, h, fpx, lh, pr), s, cons, st = render(p, labels="outside", label_pt=11,
                                           figure_width_cm=3, outfile=f"{TMP}/w2.svg")
wn = M.fit_warnings(cons, s, st, w, h, fpx, lh, pr)
(ok if wn else bad)("warns at 3 cm", f"{len(wn)} warnings")

# ------------------------------------------------- 7a. colour stability
print("== automatic colour stability ==")
_pal = M.PALETTES["spectrum"]
def _mk(spec):
    return [{"name": n, "parts": [{"type": "cds", "label": g, "size": "M",
             "dir": "fwd", "join": False, "color": ""} for g in gs]}
            for n, gs in spec]
_base = _mk([("A", ["EnzymeA", "MarkerX", "GFP"]), ("B", ["EnzymeB", "MarkerX"])])
_c0 = M.assign_palette(_base, {}, _pal, "auto")

_grown = _mk([("NEW", ["EnzymeC", "MarkerX"])]) + _base
(ok if all(M.assign_palette(_grown, {}, _pal, "auto")[k] == v
           for k, v in _c0.items()) else bad)(
    "inserting a construct keeps existing colours")

import random as _rnd
_same = True
for _seed in range(6):
    _sh = _base[:]
    _rnd.Random(_seed).shuffle(_sh)
    for _c in _sh:
        _rnd.Random(_seed).shuffle(_c["parts"])
    if M.assign_palette(_sh, {}, _pal, "auto") != _c0:
        _same = False
(ok if _same else bad)("colours do not depend on row order")

_cut = M.assign_palette(_base[:1], {}, _pal, "auto")
(ok if all(_cut[k] == _c0[k] for k in _cut) else bad)(
    "removing a construct keeps survivors' colours")

(ok if M.assign_palette(_mk([("Z", ["MarkerX"])]), {}, _pal, "auto")[("cds", "MarkerX")]
       == _c0[("cds", "MarkerX")] else bad)(
    "the same label matches across figures")

# must not depend on PYTHONHASHSEED
import subprocess as _sp
_code = ("import sys;sys.path.insert(0,%r);import cistron as M;"
         "pal=M.PALETTES['spectrum'];"
         "c=[{'name':'A','parts':[{'type':'cds','label':g,'size':'M','dir':'fwd',"
         "'join':False,'color':''} for g in ['EnzymeA','MarkerX','GFP']]}];"
         "print(M.assign_palette(c,{},pal,'auto')[('cds','MarkerX')])"
         % os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
_outs = {_sp.run([sys.executable, "-c", _code], capture_output=True,
                 text=True).stdout.strip() for _ in range(3)}
(ok if len(_outs) == 1 else bad)("stable across processes", str(_outs))

# no needless collisions while the palette has slots
_many = _mk([("A", [f"gene{i}" for i in range(12)])])
_r = M.assign_palette(_many, {}, _pal, "auto")
(ok if len(set(_r.values())) == len(_r) else bad)(
    "no collisions while the palette has room", f"{len(set(_r.values()))}/{len(_r)}")

# a pin still wins, and `auto` releases it
_pinned = M.assign_palette(_base, {"MarkerX": "#123456"}, _pal, "auto")
(ok if _pinned[("cds", "MarkerX")] == "#123456" else bad)("a colors-sheet pin wins")
(ok if M.assign_palette(_base, {"MarkerX": "#123456"}, _pal, "theme")[("cds", "MarkerX")]
       != "#123456" else bad)("palette_mode=theme overrides pins")

# ------------------------------------------- 7b. scalable vs fixed parts
print("== scalable vs fixed ==")
_sizes = ("XS", "S", "M", "L")
_scal = [t for t, v in M.PART_TYPES.items() if v.get("scale", True)]
_fix  = [t for t, v in M.PART_TYPES.items() if not v.get("scale", True)]

_bad = [t for t in _fix
        if len({M.part_width({"type": t, "size": z, "label": "X"}) for z in _sizes}) != 1]
(ok if not _bad else bad)("position markers ignore the size column",
                          f"{len(_fix)} fixed types" if not _bad else f"{_bad[:3]} varied")

_bad = []
for t in _scal:
    ws = [M.part_width({"type": t, "size": z, "label": ""}) for z in _sizes]
    if not all(a < b for a, b in zip(ws, ws[1:])):
        _bad.append(t)
(ok if not _bad else bad)("length-bearing parts grow with size",
                          f"{len(_scal)} scalable types" if not _bad else str(_bad[:3]))

# a fixed part must never stretch to fit an inside label either
_bad = [t for t in _fix
        if M.part_width({"type": t, "size": None, "label": "a very long name here"},
                        "inside", 11) != M.part_width({"type": t, "size": None, "label": ""})]
(ok if not _bad else bad)("position markers do not stretch for labels",
                          "" if not _bad else str(_bad[:3]))

# and the registry invariant: nothing scalable is drawn with an unstretchable glyph
_bad = [t for t, v in M.PART_TYPES.items()
        if v.get("scale", True) and v["glyph"] not in M.STRETCHABLE_GLYPHS]
(ok if not _bad else bad)("no scalable type uses an unstretchable glyph",
                          "" if not _bad else str(_bad))

# setting a size on a point marker must be reported, not silently dropped
_p = make_wb(f"{TMP}/scale.xlsx", [("A", 1, "loxp", "site", "L", "", ""),
                                   ("A", 2, "cds", "g", "M", "", "")])
(_o, _w, _h, _f, _l, _pr), _s, _c, _st = render(_p, figure_width_cm=16)
_wn = M.fit_warnings(_c, _s, _st, _w, _h, _f, _l, _pr)
(ok if any("position marker" in x for x in _wn) else bad)(
    "warns when size is set on a point marker")

# ------------------------------------------- 7b2. colours stay distinct
print("== colours stay distinct ==")
def _labels(n, pal="nature"):
    cons = [{"name": "A", "parts": [
        {"type": "cds", "label": f"gene{i}", "size": "M", "dir": "fwd",
         "join": False, "color": ""} for i in range(n)]}]
    return cons, M.assign_palette(cons, {}, M.PALETTES[pal], "auto")

_c, _r = _labels(14)
_v = list(_r.values())
(ok if len(set(_v)) == len(_v) else bad)(
    "14 parts on a 10-colour palette get 14 distinct colours",
    f"{len(set(_v))}/{len(_v)}")
_worst = min(M.colour_distance(_v[i], _v[j])
             for i in range(len(_v)) for j in range(i + 1, len(_v)))
(ok if _worst >= 18 else bad)("and they are visibly apart", f"min dE {_worst:.1f}")

import copy as _cp
_base, _rb = _labels(12)
_old = _rb[("cds", "gene0")]
_mod = _cp.deepcopy(_base)
for _p in _mod[0]["parts"]:
    if _p["label"] == "gene0":
        _p["label"] = "completely-different-name"
_rn = M.assign_palette(_mod, {}, M.PALETTES["nature"], "auto")
_new = _rn[("cds", "completely-different-name")]
(ok if M.colour_distance(_old, _new) >= 18 else bad)(
    "renaming a part gives a visibly different colour",
    f"dE {M.colour_distance(_old, _new):.1f}")

_ext = M.extend_palette(M.PALETTES["nature"], 30)
(ok if len(set(_ext)) == len(_ext) else bad)("extended palette has no repeats",
                                             f"{len(set(_ext))}/{len(_ext)}")
(ok if len(_ext) > len(M.PALETTES["nature"]) else bad)(
    "extend_palette actually extends",
    f"{len(M.PALETTES['nature'])} -> {len(_ext)}")

import time as _t
_big = [{"name": f"C{j}", "parts": [
    {"type": "cds", "label": f"g{j}_{i}", "size": "M", "dir": "fwd",
     "join": False, "color": ""} for i in range(8)]} for j in range(120)]
_t0 = _t.time()
M.assign_palette(_big, {}, M.PALETTES["nature"], "auto")
_dt = _t.time() - _t0
(ok if _dt < 2.0 else bad)("960 labels assign quickly", f"{_dt*1000:.0f} ms")

# ---------------------------------------------- 7c. colour vision
print("== colour vision ==")
# pairs a real dichromat confuses, cross-checked against colour-science
_BAD = [("#88551E", "#E02E2E"), ("#81E02E", "#E0E02E"), ("#5C2B7D", "#2B457D")]
_SAFE = [("#0072B2", "#E69F00"), ("#D55E00", "#009E73"),
         ("#2EADE0", "#E02E2E"), ("#CC79A7", "#009E73")]
_miss = [p for p in _BAD
         if min(M.colour_distance(*p, k) for k in M._CVD) >= M.CVD_CLASH]
_fp = [p for p in _SAFE
       if min(M.colour_distance(*p, k) for k in M._CVD) < M.CVD_CLASH]
(ok if not _miss else bad)("known colour-blind collisions are caught",
                           f"{len(_BAD)-len(_miss)}/{len(_BAD)}")
(ok if not _fp else bad)("colour-blind-safe pairs are not flagged",
                         f"{len(_SAFE)-len(_fp)}/{len(_SAFE)}")
(ok if M.colour_distance("#123456", "#123456") == 0 else bad)(
    "a colour has zero distance from itself")

# the warning must only fire where colour actually carries the identity
_cv = make_wb(f"{TMP}/cv.xlsx",
              [("A", 1, "cds", "one", "M", "", "", "#88551E"),
               ("A", 2, "cds", "two", "M", "", "", "#E02E2E")])
def _cvwarn(**over):
    (o, w_, h_, f_, l_, pr_), s_, c_, st_ = render(_cv, figure_width_cm=16, **over)
    _, pal_, _, _ = M.resolve_theme(s_)
    res_ = M.assign_palette(c_, {}, pal_, "auto")
    return any("colour-blind" in x for x in
               M.fit_warnings(c_, s_, st_, w_, h_, f_, l_, pr_, res_))
(ok if _cvwarn(labels="none", show_legend="TRUE") else bad)(
    "warns when colour is the only key")
(ok if not _cvwarn(labels="outside", show_legend="FALSE") else bad)(
    "stays quiet when every glyph is named and there is no legend")

# --------------------------------------------- 8. older-Python syntax
# Nested same-delimiter quotes and backslashes inside an f-string expression
# are legal only from Python 3.12. The code must stay importable on 3.8+.
print("== python 3.8+ syntax ==")
import re as _re
_src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                         "cistron.py")).read()
_bad = []
for _i, _line in enumerate(_src.splitlines(), 1):
    for _pat, _why in [(r"f'[^']*\{[^}]*'[^}]*\}", "nested ' in f-expression"),
                       (r'f"[^"]*\{[^}]*"[^}]*\}', 'nested " in f-expression'),
                       (r"f['\"][^'\"]*\{[^}]*\\\\", "backslash in f-expression")]:
        if _re.search(_pat, _line):
            _bad.append((_i, _why))
(ok if not _bad else bad)("no Python 3.12-only f-string syntax",
                          "" if not _bad else f"{len(_bad)} sites, first at line {_bad[0][0]}: {_bad[0][1]}")

# ------------------------------------------------------------------ report
print("\n" + "="*72)
for n, d in PASS: print(f"  PASS  {n:42} {d}")
for n, d in WARN: print(f"  WARN  {n:42} {d}")
for n, d in FAIL: print(f"  FAIL  {n:42} {d}")
print("="*72)
print(f"{len(PASS)} passed, {len(WARN)} warnings, {len(FAIL)} failed")

sys.exit(1 if FAIL else 0)
