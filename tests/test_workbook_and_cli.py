"""Workbook builder, CLI and static checks.

Run alongside test_suite.py:  python3 tests/test_workbook_and_cli.py
"""
import os, sys, io, time, tempfile, shutil, subprocess, re
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import cistron as M
from openpyxl import Workbook, load_workbook

TMP = tempfile.mkdtemp(prefix="cd_extra_")
P, F, S = [], [], []
def ok(n, d=""):   P.append((n, d))
def bad(n, d=""):  F.append((n, d))
def skip(n, d=""): S.append((n, d))

try:
    import cairosvg          # optional: only needed for .png / .pdf output
    HAVE_CAIRO = True
except Exception:
    HAVE_CAIRO = False

def wb_with(sheets):
    """sheets = {name: [row, row, ...]}"""
    wb = Workbook(); wb.remove(wb.active)
    for n, rows in sheets.items():
        ws = wb.create_sheet(n)
        for r in rows: ws.append(list(r))
    p = os.path.join(TMP, f"w{len(os.listdir(TMP))}.xlsx"); wb.save(p); return p

CAS = [("cassette","order","type","label","size","dir","join","color"),
       ("A",1,"HA","Left","M","",""),("A",2,"cds","Gene","M","","")]

# ---------------------------------------------- header detection
print("== header detection ==")
p = wb_with({"cassettes": [("Some title",),(None,),("notes about this file",),(None,)] + CAS})
try:
    _,_,cons = M.load_workbook_config(p)
    (ok if len(cons)==1 and len(cons[0]["parts"])==2 else bad)("title block above the cassettes header")
except Exception as e: bad("title block above the cassettes header", f"{type(e).__name__}: {e}")

p = wb_with({"cassettes": CAS,
             "colors": [("Colours are automatic",),(None,),("name","hex"),("cds","#112233")]})
try:
    _,c,_ = M.load_workbook_config(p)
    (ok if c.get("cds")=="#112233" else bad)("title block above the colors header", str(c.get("cds")))
except Exception as e: bad("title block above the colors header", f"{type(e).__name__}: {e}")

p = wb_with({"cassettes": CAS,
             "settings": [("How it looks",),(None,),("setting","value"),("theme","poster")]})
try:
    s,_,_ = M.load_workbook_config(p)
    (ok if s["theme"]=="poster" else bad)("title block above the settings header", str(s["theme"]))
except Exception as e: bad("title block above the settings header", f"{type(e).__name__}: {e}")

# ---------------------------------------------- auto sentinel / modes
print("== colour overrides ==")
for word in ("auto","AUTO","Automatic","-"):
    p = wb_with({"cassettes": CAS, "colors": [("name","hex"),("cds",word)]})
    _,c,_ = M.load_workbook_config(p)
    if "cds" in c: bad(f"`{word}` releases a pin", "still pinned"); break
else: ok("auto / AUTO / Automatic / - all release a pin")

p = wb_with({"cassettes": CAS, "settings":[("setting","value"),("palette_mode","sheet")]})
try:
    s,_,_ = M.load_workbook_config(p)
    (ok if s["palette_mode"]=="auto" else bad)("legacy `sheet` maps to `auto`", str(s["palette_mode"]))
except Exception as e: bad("legacy `sheet` maps to `auto`", str(e))

# ---------------------------------------------- builder round trips
print("== workbook builder ==")
p1 = os.path.join(TMP,"s1.xlsx"); M.write_starter(p1)
try:
    s,c,cons = M.load_workbook_config(p1)
    (ok if cons else bad)("a fresh --new workbook parses", f"{len(cons)} cassette(s)")
    (ok if not c else bad)("a fresh workbook pins nothing", f"{len(c)} pins")
except Exception as e: bad("a fresh --new workbook parses", str(e))

# rebuild from its own contents; must be stable
def dump(path):
    s,c,cons = M.load_workbook_config(path)
    return ([ (x["name"], tuple((q["type"],q["label"],q.get("size"),q["dir"],q["join"])
               for q in x["parts"])) for x in cons ], dict(c),
            {k:v for k,v in s.items() if not k.startswith("_")})
rows = [(x["name"] if i==0 else None, i+1, q["type"], q["label"], q.get("size"),
         q["dir"], "yes" if q["join"] else "", q.get("color") or "", "")
        for x in M.load_workbook_config(p1)[2] for i,q in enumerate(x["parts"])]
p2 = os.path.join(TMP,"s2.xlsx"); M.write_starter(p2, rows=rows, colors=[], settings=dump(p1)[2])
(ok if dump(p1)[0]==dump(p2)[0] else bad)("builder round trip keeps the constructs")

# large workbook: validation and formatting must still cover the data
print("== scale ==")
big = [(f"C{j}", i+1, "cds", f"g{j}_{i}", "M", "", "", "", "") for j in range(120) for i in range(8)]
p3 = os.path.join(TMP,"big.xlsx"); M.write_starter(p3, rows=big, colors=[], settings={"figure_width_cm":20})
w = load_workbook(p3)["cassettes"]
last_data = sum(1 for r in w.iter_rows(min_row=2,max_col=3,values_only=True) if r[2]) + 1
dv_last = max(int(re.findall(r"(\d+)$", str(d.sqref))[0]) for d in w.data_validations.dataValidation)
cf_last = max(int(re.findall(r"(\d+)$", str(k.sqref).split(":")[-1])[0])
              for k in w.conditional_formatting)
(ok if dv_last >= last_data else bad)("dropdowns reach the last data row", f"data {last_data}, dv {dv_last}")
(ok if cf_last >= last_data else bad)("conditional formatting reaches the last data row", f"data {last_data}, cf {cf_last}")
(ok if str(w.auto_filter.ref).split(":")[-1][1:].isdigit() and
    int(str(w.auto_filter.ref).split(":")[-1][1:]) >= last_data else bad)(
    "autofilter covers the data", str(w.auto_filter.ref))

t0=time.time(); s,c,cons = M.load_workbook_config(p3)
s.update(outfile=os.path.join(TMP,"big.svg"), labels="outside", label_repeat="once")
st,pal,_,_ = M.resolve_theme(s); res = M.assign_palette(cons,c,pal,"auto")
M.build_figure(cons,s,res,st); dt=time.time()-t0
(ok if dt < 25 else bad)("960 elements draw in reasonable time", f"{dt:.1f}s")

# ---------------------------------------------- CLI paths
print("== cli ==")
d = os.path.join(TMP,"cli"); os.makedirs(d)
shutil.copy(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                         "cistron.py"), d)
def run(*a, cwd=d):
    r = subprocess.run([sys.executable,"cistron.py",*a],cwd=cwd,
                       capture_output=True,text=True,timeout=120)
    return r.returncode, r.stdout+r.stderr
for args, needle, name in [
    (("--help",), "--new", "--help prints the flags"),
    (("--list",), "PALETTES", "--list prints the catalogue"),
    (("--new","x.xlsx"), "wrote", "--new writes a workbook"),
    (("x.xlsx",), "wrote", "the new workbook draws"),
    (("x.xlsx","--check"), "", "--check runs"),
    (("x.xlsx","--themes"), "samples in", "--themes writes samples"),
    (("missing.xlsx",), "No workbook found", "a missing workbook says so"),
]:
    rc,out = run(*args)
    hit = needle.lower() in out.lower()
    (ok if hit else bad)(name, out.strip().splitlines()[-1][:70] if out.strip() else f"rc={rc}")
# the missing-workbook message must hand over a command that actually works
rc, out = run("missing.xlsx")
_cmd = [l.strip() for l in out.splitlines() if l.strip().startswith("python3 ")]
if _cmd:
    _rc2, _out2 = run(*_cmd[0].split()[2:])
    (ok if _rc2 == 0 else bad)("its suggested command runs", _cmd[0])
else:
    bad("its suggested command runs", "no command offered")

# in an empty folder it should point at --new, not at a file that is not there
_empty = os.path.join(TMP, "empty"); os.makedirs(_empty, exist_ok=True)
shutil.copy(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                         "cistron.py"), _empty)
_rc3, _out3 = run(cwd=_empty)
(ok if "--new" in _out3 else bad)("an empty folder is told how to start",
                                  _out3.strip().splitlines()[-1][:60] if _out3.strip() else "")

# --check must not leave a figure behind
rc,out = run("x.xlsx","--check")
(ok if not any(f.endswith(".svg") and f.startswith(".cassette") for f in os.listdir(d)) else bad)(
    "--check leaves no stray file")

# png / pdf export - cairosvg is optional, so this is a skip, not a failure
for ext in (("png","pdf") if HAVE_CAIRO else ()):
    wbp = os.path.join(d,"x.xlsx"); wx = load_workbook(wbp)
    ws = wx["settings"]
    for r in range(1, ws.max_row+1):
        if str(ws.cell(row=r,column=1).value or "").strip()=="outfile":
            ws.cell(row=r,column=2,value=f"out.{ext}")
    wx.save(wbp)
    rc,out = run("x.xlsx")
    (ok if os.path.exists(os.path.join(d,f"out.{ext}")) else bad)(f"{ext} export", out[-70:])

# ---------------------------------------------- awkward content
print("== awkward content ==")
cases = {
 "label that is only spaces":  [("A",1,"cds","   ","M","",""),("A",2,"HA","R","M","","")],
 "numeric cassette name":      [(2024,1,"cds","G","M","",""),(2024,2,"HA","R","M","","")],
 "numeric label":              [("A",1,"cds",42,"M","",""),("A",2,"HA","R","M","","")],
 "same label, two types":      [("A",1,"cds","X","M","",""),("A",2,"promoter","X","","","")],
 "duplicate element rows":     [("A",1,"cds","G","M","",""),("A",1,"cds","G","M","","")],
 "very many parts in one row": [("A",i+1,"cds",f"g{i}","XS","","") for i in range(60)],
}
for name, rows in cases.items():
    p = wb_with({"cassettes":[CAS[0]]+rows})
    try:
        s,c,cons = M.load_workbook_config(p)
        s.update(outfile=os.path.join(TMP,"a.svg"), labels="outside", figure_width_cm=20)
        st,pal,_,_ = M.resolve_theme(s); res = M.assign_palette(cons,c,pal,"auto")
        out,w,h,*_ = M.build_figure(cons,s,res,st)
        import xml.etree.ElementTree as ET; ET.fromstring(open(out).read())
        ok(name, f"{w:.0f}x{h:.0f}")
    except Exception as e: bad(name, f"{type(e).__name__}: {e}")

# ---------------------------------------------- static hygiene
print("== static ==")
src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                        "cistron.py")).read()
(ok if "def main(" in src and src.count("def main(")==1 else bad)("one main()")
muts = re.findall(r"def \w+\([^)]*=\s*(\[\]|\{\})", src)
(ok if not muts else bad)("no mutable default arguments", str(muts))
bare = re.findall(r"except\s*:", src)
(ok if not bare else bad)("no bare except", f"{len(bare)}")
(ok if "import pdb" not in src and "breakpoint()" not in src else bad)("no debugger left in")
import ast as _ast
_fns = [(n.name, n.end_lineno - n.lineno + 1) for n in _ast.walk(_ast.parse(src))
        if isinstance(n, _ast.FunctionDef)]
_worst = max(_fns, key=lambda x: x[1]) if _fns else ("", 0)
(ok if _worst[1] < 200 else bad)("no function over 200 lines",
                                 f"longest: {_worst[0]} at {_worst[1]}")
# and the module should not be one giant wall either
(ok if len(_fns) > 40 else bad)("work is split across functions", f"{len(_fns)} functions")
print()
if not HAVE_CAIRO:
    skip("png / pdf export", "cairosvg not installed (optional)")
    # the refusal path must still be helpful
    d2 = os.path.join(TMP, "nc"); os.makedirs(d2, exist_ok=True)
    shutil.copy(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                             "cistron.py"), d2)
    subprocess.run([sys.executable, "cistron.py", "--new", "n.xlsx"],
                   cwd=d2, capture_output=True, text=True)
    wx = load_workbook(os.path.join(d2, "n.xlsx")); wsx = wx["settings"]
    for rr in range(1, wsx.max_row + 1):
        if str(wsx.cell(row=rr, column=1).value or "").strip() == "outfile":
            wsx.cell(row=rr, column=2, value="out.png")
    wx.save(os.path.join(d2, "n.xlsx"))
    rc2 = subprocess.run([sys.executable, "cistron.py", "n.xlsx"],
                         cwd=d2, capture_output=True, text=True)
    out2 = rc2.stdout + rc2.stderr
    (ok if ("cairosvg" in out2 and "pip install" in out2
            and os.path.exists(os.path.join(d2, "out.svg"))) else bad)(
        "missing cairosvg explains itself and still writes the SVG",
        out2.strip().splitlines()[-1][:60] if out2.strip() else "")

for n,d in P: print(f"  PASS  {n:46} {d}")
for n,d in S: print(f"  SKIP  {n:46} {d}")
for n,d in F: print(f"  FAIL  {n:46} {d}")
print(f"\n{len(P)} passed, {len(S)} skipped, {len(F)} failed")
shutil.rmtree(TMP, ignore_errors=True)

sys.exit(1 if F else 0)
