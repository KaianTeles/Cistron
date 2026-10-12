#!/usr/bin/env python3
"""Regressions from the 1.0 review.

Every check here failed before the fix it guards. They are kept apart from
the main suite so it stays readable as a description of the tool, and this
stays readable as a list of mistakes not to make twice.
"""
import itertools
import os
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import cistron as M                                              # noqa: E402

PASS = FAIL = 0


def ok(name, note=""):
    global PASS
    PASS += 1
    print(f"  PASS  {name:52} {note}")


def bad(name, note=""):
    global FAIL
    FAIL += 1
    print(f"  FAIL  {name:52} {note}")


def check(cond, name, note=""):
    (ok if cond else bad)(name, note)


def colours(labels, pal, colors=None, mode="auto"):
    cons = [{"name": "c", "parts": [
        {"type": "cds", "label": l, "size": "M", "dir": "fwd",
         "join": False, "color": ""} for l in labels]}]
    return M.assign_palette(cons, colors or {}, list(M.PALETTES[pal]), mode)


def only(labels, pal):
    r = colours(labels, pal)
    return {l: r[("cds", l)] for l in labels}


# ------------------------------------------------- colour stays put
print("== automatic colour is stable under edits ==")
# Measured over many random label sets, because any single set can be an
# unlucky draw. The claim is about the typical figure, so the test is too.
import random as _rnd                                            # noqa: E402
import statistics as _stat                                       # noqa: E402
_r = _rnd.Random(11)
for pal in ("okabe-ito", "nature", "spectrum", "tol-muted", "grayscale"):
    for n in (8, 14, 25):
        moved = []
        for _ in range(15):
            L = [f"p{_r.randrange(10 ** 6)}" for _ in range(n)]
            a = only(L, pal)
            b = only(L + [f"x{_r.randrange(10 ** 6)}"], pal)
            moved.append(sum(a[l] != b[l] for l in L))
        med = _stat.median(moved)
        # The bound is loose on purpose. Keeping every colour in a figure
        # visibly distinct is a property of the whole set, so a newcomer
        # that forces a rearrangement moves others with it. What must hold
        # is that it is a minority, not that it is nobody.
        check(med <= n * 0.4,
              f"adding a part to {pal} leaves most of {n} colours alone",
              f"median {med:.0f} moved, worst {max(moved)}")

L = ["cas9", "gfp", "rfp", "hygR", "AmpR", "pTEF1", "tNOS", "loxP"]
import random                                                    # noqa: E402
rnd = random.Random(7)
same = all(colours(L, "nature") == colours(rnd.sample(L, len(L)), "nature")
           for _ in range(20))
check(same, "reordering rows changes nothing at all")

before = only(L, "nature")
after = only(L[:-1], "nature")
_moved = sum(before[l] != after[l] for l in L[:-1])
check(_moved <= len(L) * 0.4, "removing a part leaves most colours alone",
      f"{_moved} of {len(L) - 1} moved")

check(only(["a", "b", "c"], "nature") == only(["a", "b", "c"], "nature"),
      "the same figure twice gives the same colours")

# ------------------------------------------------- no two parts alike
print("== every part gets its own colour ==")
for pal in sorted(M.PALETTES):
    for n in (9, 16, 30):
        v = list(only([f"g{i:02d}" for i in range(n)], pal).values())
        if len(set(v)) != len(v):
            bad(f"{pal} draws {n} parts in distinct colours",
                f"{len(set(v))}/{n}")
            break
    else:
        ok(f"{pal} draws 9, 16 and 30 parts in distinct colours")

v = list(only([f"gene{i}" for i in range(14)], "nature").values())
worst = min(M.colour_distance(a, b) for a, b in itertools.combinations(v, 2))
check(worst >= M.SAME_COLOUR, "14 parts on a 10-colour palette stay apart",
      f"min dE {worst:.1f}")

# a pinned colour is never approached by an automatic one
print("== a pin keeps its colour to itself ==")
_pin = "#56B4E9"
_r = colours([f"q{i}" for i in range(18)], "okabe-ito", {"q0": _pin})
_near = min(M.colour_distance(v, _pin) for k, v in _r.items() if k[1] != "q0")
check(_near >= M.SAME_COLOUR, "no automatic colour lands on a pinned one",
      f"closest dE {_near:.1f}")

# and the figure holds up for a colour-blind reader where the palette allows
print("== colours survive red-green colour blindness ==")
for _pal in ("okabe-ito", "tol-bright", "tol-muted", "spectrum"):
    _v = list(only([f"g{i}" for i in range(8)], _pal).values())
    _w = min(M.colour_distance(a, b, k)
             for a, b in itertools.combinations(_v, 2)
             for k in ("protanopia", "deuteranopia"))
    check(_w >= M.CVD_CLASH, f"8 parts on {_pal} stay apart for a CVD reader",
          f"min dE {_w:.1f}")

check(len(M.extend_palette(list(M.PALETTES["grayscale"]), 12)) == 12,
      "grayscale can be extended at all")

# ------------------------------------------------- advice that is true
print("== the sanity check gives advice that works ==")
for pal in ("okabe-ito", "tol-bright", "tol-muted"):
    cs = list(M.PALETTES[pal])
    clash = [(a, b) for a, b in itertools.combinations(cs, 2)
             if min(M.colour_distance(a, b, k) for k in M._CVD) < M.CVD_CLASH]
    check(not clash, f"{pal}, which the CVD warning recommends, is safe",
          f"{len(clash)} clashing pairs")

src = open(os.path.join(ROOT, "cistron.py"), encoding="utf-8").read()
check("or cividis " not in src,
      "cividis is not recommended as a colour-blind-safe palette")

# ------------------------------------------------- nothing drawn off-canvas
print("== nothing is drawn outside the canvas ==")
import re                                                        # noqa: E402

TMP = tempfile.mkdtemp()


def render(cons, **kw):
    s = dict(M.SETTINGS)
    s.update(_basedir=TMP, outfile="t.svg", **kw)
    st, pal, _, _ = M.resolve_theme(s)
    r = M.assign_palette(cons, {}, pal, "theme")
    path, w, h = M.build_figure(cons, s, r, st)[:3]
    return open(path, encoding="utf-8").read(), w, h


def mk(name, parts):
    return {"name": name, "parts": [
        {"type": M.ALIASES[t], "label": l, "size": sz, "dir": "fwd",
         "join": False, "color": ""} for t, l, sz in parts]}


def escapes(svg, w):
    g = re.search(r"translate\(([\d.]+),0\)", svg)
    shift = float(g.group(1)) if g else 0.0
    out = []
    # parse the attributes properly: they are not emitted in a fixed order,
    # and guessing it is how this check quietly passed everything once
    for tag, txt in re.findall(r"<text([^>]*)>([^<]*)</text>", svg):
        at = dict(re.findall(r'([\w-]+)="([^"]*)"', tag))
        x, fs = float(at["x"]), float(at["font-size"])
        anchor = at.get("text-anchor", "start")
        tw = M.text_width(txt, fs)
        left = {"middle": x - tw / 2, "end": x - tw}.get(anchor, x) + shift
        if left < -0.5 or left + tw > w + 0.5:
            out.append((txt[:18], round(left, 1), round(left + tw, 1)))
    return out


LONG = "pXX-reporter-overexpression-line-7-verified-clone-B2"
cases = [
    ("a long name above the row",
     [mk(LONG, [("prom", "P", None), ("cds", "cbs", "XS")])],
     dict(labels="none", row_titles="above", show_legend="FALSE")),
    ("a long name beside the row",
     [mk(LONG, [("prom", "P", None), ("cds", "cbs", "XS")])],
     dict(labels="none", row_titles="left", show_legend="FALSE")),
    ("wide capitals in the legend",
     [mk("r", [("cds", l, "L") for l in
               ("WWWWWWWWWWWWWWWWWWWW", "MMMMMMMMMMMMMMMMMMMM", "short")])],
     dict(labels="none", show_legend="TRUE", figure_width_cm=None)),
    ("a long label on the last part",
     [mk("r", [("prom", "P", None), ("loxp", "A very long trailing name", None)])],
     dict(labels="outside", show_legend="FALSE")),
    ("a long label on the first part",
     [mk("r", [("loxp", "A very long leading name", None), ("prom", "P", None)])],
     dict(labels="outside", show_legend="FALSE")),
]
for name, cons, kw in cases:
    for side in ("left", "right"):
        svg, w, _ = render(cons, anchor_side=side, **kw)
        esc = escapes(svg, w)
        if esc:
            bad(f"{name} ({side})", str(esc[:2]))
            break
    else:
        ok(name)

# ------------------------------------------------- CLI safety
print("== the CLI does not destroy anything ==")
CLI = os.path.join(ROOT, "cistron.py")


def run(cwd, *a):
    r = subprocess.run([sys.executable, CLI, *a], cwd=cwd,
                       capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


d = tempfile.mkdtemp()
run(d, "--new", "keep.xlsx")
before = open(os.path.join(d, "keep.xlsx"), "rb").read()
rc, out = run(d, "--new", "keep.xlsx")
after = open(os.path.join(d, "keep.xlsx"), "rb").read()
check(rc != 0 and before == after, "--new refuses to overwrite a workbook",
      out.strip().splitlines()[0][:48] if out.strip() else "")

rc, out = run(d, "keep.xlsx")
check(rc == 0, "a workbook made by --new is found in the same folder",
      out.strip().splitlines()[-1][:48])

with open(os.path.join(d, "broken.xlsx"), "w") as fh:
    fh.write("not a zip")
rc, out = run(d, "broken.xlsx")
check(rc != 0 and "Traceback" not in out,
      "a corrupt workbook says so instead of crashing",
      out.strip().splitlines()[0][:48])

# ------------------------------------------------- labels stay readable
print("== leader lines do not strike through labels ==")
_c = [mk("r", [("cds", ch * 22, "XS") for ch in "ABC"])]
svg, _w, _h = render(_c, labels="outside", show_legend="FALSE")
check("<line x1=" in svg and svg.index("<line x1=") < svg.index("<text x="),
      "every leader is drawn before every label")
check(svg.count("<text") == 2 * len([1 for _ in "ABC"]),
      "each label is drawn twice: a halo and the text",
      f"{svg.count('<text')} text elements")
# cairosvg - which is what this tool uses to write .png and .pdf - ignores
# paint-order and paints the outline over the glyphs, erasing the label
check("paint-order" not in svg,
      "no reliance on paint-order, which cairosvg ignores")
try:
    import cairosvg
    from PIL import Image
    import io as _io
    png = cairosvg.svg2png(bytestring=svg.encode(), output_width=1200)
    g = Image.open(_io.BytesIO(png)).convert("L")
    dark = sum(1 for px in g.convert('L').tobytes() if px < 128)
    check(dark > 2000, "labels survive a cairosvg render", f"{dark} dark pixels")
except ImportError:
    print("  SKIP  cairosvg not installed")

# ------------------------------------------------- palettes behave
print("== the shipped palettes do not fight the sanity check ==")
for name in sorted(M.PALETTES):
    if name in M.ORDERED_PALETTES:
        continue
    cs = list(M.PALETTES[name])
    clash = [(a, b) for a, b in itertools.combinations(cs, 2)
             if M.colour_distance(a, b) < M.SAME_COLOUR]
    check(not clash, f"{name} holds no two colours that read as one",
          str(clash[:2]))
check(M.ORDERED_PALETTES <= set(M.PALETTES),
      "every palette named as an ordered ramp exists")

# ------------------------------------------------- CLI refuses nonsense
print("== the CLI refuses what it cannot do ==")
d2 = tempfile.mkdtemp()
run(d2, "--new", "a.xlsx")
before = sorted(os.listdir(d2))
rc, out = run(d2, "a.xlsx", "--themes", "--check")
check(rc != 0 and sorted(os.listdir(d2)) == before,
      "--themes and --check together are refused, writing nothing")
rc, out = run(d2, "a.xlsx", "--chek")
check(rc != 0 and "Unknown option" in out,
      "a mistyped flag is refused, not ignored")
rc, out = run(d2, "-V")
check(rc == 0 and "cistron" in out, "-V prints the version", out.strip()[:30])

# ------------------------------------------------- the starter workbook
print("== a brand new workbook is valid on its own terms ==")
import openpyxl                                                  # noqa: E402
_wb = openpyxl.load_workbook(os.path.join(d2, "a.xlsx"))
_types = [r[2] for r in _wb["cassettes"].iter_rows(min_row=2, values_only=True)
          if r[2]]
check(all(t in M.PART_TYPES for t in _types),
      "its demo rows use types its own dropdown accepts", ", ".join(_types[:3]))
_over = [str(dv.sqref) for dv in _wb["settings"].data_validations.dataValidation
         if len(dv.formula1) > 255]
check(not _over, "no dropdown exceeds Excel's 255-character limit", str(_over))

print("== labels survive the round trip through Excel ==")
_rows = [("C1", i, "cds", lab, "M", "", "", "")
         for i, lab in enumerate(["=SUM(A1:A2)", "+plus", "-dash", "@at",
                                  "1-10", "TRUE"], 1)]
_p = os.path.join(tempfile.mkdtemp(), "rt.xlsx")
M.write_starter(_p, rows=_rows)
_back = [p["label"] for p in M.load_workbook_config(_p)[2][0]["parts"]]
check(_back == [r[3] for r in _rows],
      "a label that looks like a formula comes back intact", str(_back[:2]))

_p2 = os.path.join(tempfile.mkdtemp(), "hex.xlsx")
try:
    M.write_starter(_p2, colors=[("GFP", "#28B", "")])
    ok("a three-digit hex does not break the workbook writer")
except Exception as exc:
    bad("a three-digit hex does not break the workbook writer",
        f"{type(exc).__name__}: {exc}")

# ------------------------------------------------- nothing destroys a file
print("== the tool never writes over something that is not a figure ==")
d3 = tempfile.mkdtemp()
run(d3, "--new", "book.xlsx")
_wbp = os.path.join(d3, "book.xlsx")


def _set(path, key, value):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["settings"]
    for row in ws.iter_rows():
        for c in row:
            if c.value == key:
                ws.cell(c.row, c.column + 1, value)
    wb.save(path)


for value, why in [("book.xlsx", "the workbook itself"),
                   ("notes.txt", "a text file"),
                   ("../out.svg", "a path outside the folder")]:
    _set(_wbp, "outfile", value)
    size = os.path.getsize(_wbp)
    rc, out = run(d3, "book.xlsx")
    same = os.path.getsize(_wbp) == size
    check(rc != 0 and same and "wrote" not in out,
          f"outfile = {value} is refused before anything is written",
          "" if same else "FILE WAS MODIFIED")
_set(_wbp, "outfile", "fig.svg")

for value in ("nan", "inf", "-3"):
    _set(_wbp, "figure_width_cm", value)
    rc, out = run(d3, "book.xlsx")
    check(rc != 0 and "Traceback" not in out,
          f"figure_width_cm = {value} is refused with a message",
          out.strip().splitlines()[0][:46] if out.strip() else "")
_set(_wbp, "figure_width_cm", 16)

rc, out = run(d3, "book.xlsx")
_fig = os.path.join(d3, "fig.svg")
_good = os.path.getsize(_fig)
_r = subprocess.run(["bash", "-c", f"ulimit -f 1; exec '{sys.executable}' '{CLI}' book.xlsx"],
                    cwd=d3, capture_output=True, text=True)
check(os.path.getsize(_fig) == _good and "Traceback" not in _r.stdout + _r.stderr,
      "a write that fails leaves the previous figure intact",
      f"{_good} -> {os.path.getsize(_fig)} bytes")

# ------------------------------------------------- the key matches the figure
print("== the legend draws what the figure draws ==")
for _t in ("loxp", "ori", "promoter", "restriction_site", "rbs", "scar"):
    _p = {"type": _t, "label": _t, "size": None, "dir": "fwd",
          "join": False, "color": ""}
    _fig_w = M.part_width(_p, "none", 11.0)
    _leg = dict(_p, size=M.LEGEND_GLYPH)
    _leg_w = (M.SIZE_W[M.LEGEND_GLYPH] if M.is_scalable(_p)
              else M.part_width(_leg, "none", 11.0))
    check(abs(_fig_w - _leg_w) < 0.5,
          f"a {_t} is the same width in the key as on the backbone",
          f"{_fig_w:.0f} vs {_leg_w:.0f}")

print("== a reversed glyph is a mirror of the forward one ==")
_st = dict(M.resolve_style("classic"))
_st["_sw"] = 2.0


def _spans(d):
    return sorted((float(a), float(b)) for a, b in
                  re.findall(r"M ([\d.]+),[\d.]+ H ([\d.]+)", d))


_f = _spans(M._step(0, 50, 40, "#111", _st, "fwd"))
_r2 = _spans(M._step(0, 50, 40, "#111", _st, "rev"))
check(sorted((40 - b, 40 - a) for a, b in _f) == _r2,
      "overhang reversed is the forward glyph mirrored", f"{_f} vs {_r2}")

print("== the sanity check gives advice that clears it ==")
_c2 = [mk("r", [("cds", "GFP", "M")])]
_s2 = dict(M.SETTINGS)
_s2.update(_basedir=TMP, outfile="t.svg", labels="outside", cassette_spacing=20)
_st2, _pal2, _, _ = M.resolve_theme(_s2)
_r3 = M.assign_palette(_c2, {}, _pal2, "theme")
_fg = M.build_figure(_c2, _s2, _r3, _st2, write=False)
_wn = M.fit_warnings(_c2, _s2, _st2, _fg.width, _fg.height, _fg.legend_px,
                     _fg.legend_h, _fg.pitch_ratio, _r3,
                     row_pitch=_fg.row_pitch, tiers=_fg.tiers)
check(not any("stacked" in w for w in _wn),
      "the tier warning stays quiet when nothing is stacked",
      f"tiers={_fg.tiers}")


def _warn_at(width, parts):
    s = dict(M.SETTINGS)
    s.update(_basedir=TMP, outfile="t.svg", labels="outside",
             figure_width_cm=width)
    st_, pal_, _, _ = M.resolve_theme(s)
    r_ = M.assign_palette(parts, {}, pal_, "auto")
    f_ = M.build_figure(parts, s, r_, st_, write=False)
    return M.fit_warnings(parts, s, st_, f_.width, f_.height, f_.legend_px,
                          f_.legend_h, f_.pitch_ratio, r_,
                          row_pitch=f_.row_pitch, tiers=f_.tiers)


_tried = _cleared = 0
for _k in (2, 3, 4, 5):
    _cons = [mk("r", [("cds", f"Gene {i}", "M") for i in range(_k)])]
    for _w0 in (2.5, 3.0, 4.0, 5.0):
        _hit = [re.search(r"widen figure_width_cm to ([\d.]+) cm", w)
                for w in _warn_at(_w0, _cons)]
        _hit = [h for h in _hit if h]
        if not _hit:
            continue
        _tried += 1
        if not any("under your min_label_pt" in w
                   for w in _warn_at(float(_hit[0].group(1)), _cons)):
            _cleared += 1
check(_tried and _cleared == _tried,
      "following the advised width actually clears the warning",
      f"{_cleared} of {_tried}")

print("== a control character cannot produce an unopenable file ==")
_svg, _, _ = render([mk("r", [("cds", "ctl\x0bhere\x00x", "M")])],
                    labels="outside", show_legend="FALSE")
try:
    import xml.etree.ElementTree as _ET
    _ET.fromstring(_svg)
    ok("a label with a control character still parses")
except Exception as _exc:
    bad("a label with a control character still parses", str(_exc)[:40])

# ------------------------------------------------- dependencies are stated
print("== every dependency is written down ==")
import ast as _ast                                               # noqa: E402

_req = open(os.path.join(ROOT, "requirements.txt"), encoding="utf-8").read().lower()
_std = set(sys.stdlib_module_names)
_third = {}
for _root, _dirs, _files in os.walk(ROOT):
    _dirs[:] = [d for d in _dirs if d not in (".git", "__pycache__")]
    for _f in _files:
        if not _f.endswith(".py"):
            continue
        _p = os.path.join(_root, _f)
        for _n in _ast.walk(_ast.parse(open(_p, encoding="utf-8").read())):
            _mods = []
            if isinstance(_n, _ast.Import):
                _mods = [a.name.split(".")[0] for a in _n.names]
            elif isinstance(_n, _ast.ImportFrom) and _n.level == 0 and _n.module:
                _mods = [_n.module.split(".")[0]]
            for _m in _mods:
                if _m not in _std and _m != "cistron":
                    _third.setdefault(_m, set()).add(os.path.relpath(_p, ROOT))

# what pip is asked to install, not what the module is called
_PIP = {"PIL": "pillow"}
for _m, _where in sorted(_third.items()):
    _name = _PIP.get(_m, _m).lower()
    check(_name in _req, f"{_m} is named in requirements.txt",
          f"imported by {sorted(_where)[0]}")
check(all(_m in ("openpyxl", "cairosvg", "PIL") for _m in _third),
      "no dependency has crept in unnoticed", ", ".join(sorted(_third)))

# the one hard dependency must be the only one the tool itself needs
_core = set()
for _n in _ast.walk(_ast.parse(open(os.path.join(ROOT, "cistron.py"),
                                    encoding="utf-8").read())):
    if isinstance(_n, _ast.Import):
        _core |= {a.name.split(".")[0] for a in _n.names}
    elif isinstance(_n, _ast.ImportFrom) and _n.level == 0 and _n.module:
        _core.add(_n.module.split(".")[0])
check(_core - _std <= {"openpyxl", "cairosvg"},
      "the tool itself imports nothing beyond openpyxl and cairosvg",
      ", ".join(sorted(_core - _std)))

# ------------------------------------------------- the CLI takes one job
print("== the command line refuses what it cannot honour ==")
d4 = tempfile.mkdtemp()
rc, out = run(d4, "--new", "x.xlsx", "--check")
check(rc != 0 and not os.path.exists(os.path.join(d4, "x.xlsx")),
      "--new with --check is refused, and writes nothing")
run(d4, "--new", "x.xlsx")
rc, out = run(d4, "-l")
check(rc == 0 and "palette" in out.lower(), "-l still lists the catalogue",
      out.strip().splitlines()[0][:40] if out.strip() else f"rc={rc}")
rc, out = run(d4, "x.xlsx", "y.xlsx")
check(rc != 0 and "One workbook" in out,
      "two workbooks at once is refused rather than half-ignored")

print("== a misspelled setting is reported, not swallowed ==")
_d6 = tempfile.mkdtemp()
run(_d6, "--new", "m.xlsx")
import openpyxl as _op                                           # noqa: E402
_w6 = _op.load_workbook(os.path.join(_d6, "m.xlsx"))
for _row in _w6["settings"].iter_rows():
    for _c in _row:
        if _c.value == "figure_width_cm":
            _c.value = "figure_with_cm"
_w6.save(os.path.join(_d6, "m.xlsx"))
_rc, _out = run(_d6, "m.xlsx")
check(_rc == 0 and "no such setting" in _out,
      "a misspelled setting name is reported", _out.strip().splitlines()[0][:44])

print("== a setting that looks like a formula survives ==")
_p5 = os.path.join(tempfile.mkdtemp(), "s.xlsx")
M.write_starter(_p5, settings={"outfile": "=figure.svg"})
_s5 = M.load_workbook_config(_p5)[0]
check(_s5["outfile"] == "=figure.svg",
      "outfile written as =figure.svg reads back unchanged",
      repr(_s5["outfile"]))

# ------------------------------------------------- one way to install
print("== the install instructions do not contradict themselves ==")
_rdme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
_blocks = re.findall(r"```bash\n(.*?)```", _rdme, re.S)
check(_blocks, "the README has a bash block")
if _blocks:
    check("requirements.txt" in _blocks[0],
          "the first command block installs from requirements.txt",
          _blocks[0].strip().splitlines()[0][:40])
    _installers = [b for b in _blocks if "pip install" in b]
    _recipes = {tuple(l.strip() for l in b.strip().splitlines()
                      if l.strip().startswith(("pip ", "git ", "cd ")))
                for b in _installers}
    check(len(_recipes) <= 1,
          "there is only one install recipe, not two that differ",
          f"{len(_recipes)} different ones")
check("pip install -r requirements.txt" in _rdme,
      "it points at requirements.txt, so the dependency list has one home")

# ------------------------------------------------- the gallery is legible
print("== the gallery panels are built to be comparable ==")
_gal = open(os.path.join(ROOT, "examples", "make_gallery.py"),
            encoding="utf-8").read()
check("_wrap(" in _gal and "caption too wide" in _gal,
      "captions are wrapped to the cell and asserted, not drawn blind")
check("output_width=max(1, round(fig.width * SCALE))" in _gal,
      "panels are rendered at a fixed scale per drawing unit",
      "so two panels side by side can be compared")

# ------------------------------------------------- the note is findable
print("== known issues are written down and easy to find ==")
_ki = os.path.join(ROOT, "KNOWN_ISSUES.md")
check(os.path.exists(_ki), "KNOWN_ISSUES.md exists")
_kitxt = open(_ki, encoding="utf-8").read() if os.path.exists(_ki) else ""
_rd = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
check("KNOWN_ISSUES.md" in _rd, "the README links to it")
check(_rd.index("KNOWN_ISSUES.md") < len(_rd) * 0.15,
      "the link is near the top of the README, not buried")
check("KNOWN_ISSUES.md" in src, "--help points at it")
# each entry has to be reproducible, so each has to name its threshold
check(_kitxt.count("###") >= 5, "it lists the issues separately",
      f"{_kitxt.count('###')} entries")
check("Measured" in _kitxt or "measured" in _kitxt,
      "it says these were measured, not guessed")

# ------------------------------------------------- the README keeps up
print("== the README names everything the tool offers ==")
readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read().lower()


def documented(title, names):
    missing = [n for n in names if str(n).lower() not in readme]
    check(not missing, title, ", ".join(map(str, missing[:6])) or "")


documented("every setting is in the README",
           [k for k in M.SETTINGS if not k.startswith("_")])
documented("every theme is in the README", sorted(M.THEMES))
documented("every palette is in the README", sorted(M.PALETTES))
documented("every style is in the README", sorted(M.STYLES))
documented("every part type is in the README", sorted(M.PART_TYPES))
documented("every size key is in the README", sorted(M.SIZE_W))
documented("every command-line flag is in the README",
           sorted(set(re.findall(r'"(--[a-z-]+)"', src))))

print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
